"""Generate a synthetic vendor DOCX with seeded PII canary values."""

from __future__ import annotations

import tempfile
from pathlib import Path

from tests.fixtures.make_pdf import CANARIES


def make_synthetic_docx(output_dir: Path | None = None) -> Path:
    """Generate a synthetic vendor DOCX and return its path."""
    try:
        from docx import Document  # type: ignore[import-untyped]
        from docx.shared import Pt  # type: ignore[import-untyped]
    except ImportError as exc:
        raise ImportError("python-docx is required: pip install python-docx") from exc

    if output_dir is None:
        output_dir = Path(tempfile.mkdtemp())

    output_path = output_dir / "synthetic_vendor_dpa.docx"
    doc = Document()

    doc.add_heading("Data Processing Agreement", level=1)
    doc.add_paragraph("This Data Processing Agreement is entered into between the parties.")

    doc.add_heading("1. Data Controller", level=2)
    p = doc.add_paragraph()
    p.add_run(f"Authorized Representative: {CANARIES['PERSON']}")
    p = doc.add_paragraph()
    p.add_run(f"Email: {CANARIES['EMAIL']}")
    p = doc.add_paragraph()
    p.add_run(f"Phone: {CANARIES['PHONE']}")

    doc.add_heading("2. Personal Data Categories", level=2)
    doc.add_paragraph(
        "The processor shall process the following categories of personal data: "
        "name, email address, phone number, national identification numbers."
    )

    # Table with PII
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "SSN"
    table.cell(0, 2).text = "Payment Card"
    table.cell(1, 0).text = CANARIES["PERSON"]
    table.cell(1, 1).text = CANARIES["SSN"]
    table.cell(1, 2).text = CANARIES["CREDIT_CARD"]

    doc.add_heading("3. Security Measures", level=2)
    doc.add_paragraph(
        f"The processor maintains encryption keys in AWS KMS. "
        f"API credentials: {CANARIES['API_KEY']}. "
        f"Infrastructure hosted at {CANARIES['HOSTNAME']}."
    )

    doc.add_heading("4. IBAN for Payments", level=2)
    doc.add_paragraph(f"Bank Account (IBAN): {CANARIES['IBAN']}")

    doc.save(str(output_path))
    return output_path


if __name__ == "__main__":
    p = make_synthetic_docx(Path("."))
    print(f"Generated: {p}")
