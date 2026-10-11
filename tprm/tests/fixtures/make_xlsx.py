"""Generate a synthetic vendor XLSX with seeded PII canary values."""

from __future__ import annotations

import tempfile
from pathlib import Path

from tests.fixtures.make_pdf import CANARIES


def make_synthetic_xlsx(output_dir: Path | None = None) -> Path:
    """Generate a synthetic vendor XLSX and return its path."""
    try:
        import openpyxl  # type: ignore[import-untyped]
        from openpyxl.styles import Font  # type: ignore[import-untyped]
    except ImportError as exc:
        raise ImportError("openpyxl is required: pip install openpyxl") from exc

    if output_dir is None:
        output_dir = Path(tempfile.mkdtemp())

    output_path = output_dir / "synthetic_vendor_questionnaire.xlsx"
    wb = openpyxl.Workbook()

    # ── Sheet 1: Vendor Contact Info ─────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "VendorContacts"

    ws1["A1"] = "Field"
    ws1["B1"] = "Value"
    ws1["A2"] = "Primary Contact"
    ws1["B2"] = CANARIES["PERSON"]
    ws1["A3"] = "Email"
    ws1["B3"] = CANARIES["EMAIL"]
    ws1["A4"] = "Phone"
    ws1["B4"] = CANARIES["PHONE"]
    ws1["A5"] = "SSN (Tax ID)"
    ws1["B5"] = CANARIES["SSN"]
    ws1["A6"] = "Bank IBAN"
    ws1["B6"] = CANARIES["IBAN"]
    ws1["A7"] = "Credit Card"
    ws1["B7"] = CANARIES["CREDIT_CARD"]

    # ── Sheet 2: Security Controls ────────────────────────────────────────────
    ws2 = wb.create_sheet("SecurityControls")
    ws2["A1"] = "Control"
    ws2["B1"] = "Status"
    ws2["C1"] = "Evidence"
    ws2["A2"] = "Multi-Factor Authentication"
    ws2["B2"] = "Implemented"
    ws2["C2"] = "All users required to use MFA"
    ws2["A3"] = "Encryption at Rest"
    ws2["B3"] = "Implemented"
    ws2["C3"] = f"AES-256; KMS key ID: {CANARIES['API_KEY']}"
    ws2["A4"] = "Network Segmentation"
    ws2["B4"] = "Partial"
    ws2["C4"] = f"Primary segment: {CANARIES['IP']}/24"
    ws2["A5"] = "Database Host"
    ws2["B5"] = "Configured"
    ws2["C5"] = CANARIES["HOSTNAME"]

    # ── Sheet 3: Hidden sheet with PII ────────────────────────────────────────
    ws3 = wb.create_sheet("InternalNotes")
    ws3.sheet_state = "hidden"
    ws3["A1"] = "Internal Reference"
    ws3["B1"] = "Do not share"
    ws3["A2"] = "Employee SSN"
    ws3["B2"] = CANARIES["SSN"]
    ws3["A3"] = "Legacy API Key"
    ws3["B3"] = CANARIES["API_KEY"]

    wb.save(str(output_path))
    return output_path


if __name__ == "__main__":
    p = make_synthetic_xlsx(Path("."))
    print(f"Generated: {p}")
